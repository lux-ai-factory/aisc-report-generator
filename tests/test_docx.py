"""DOCX export.

mode=docx builds the PDF's document HTML and converts it with report_renderer/docx.py (python-docx).
Read here with python-docx (dev dependency) and the raw package XML.
"""
import re

import pytest

from conftest import need
from v2_core_helpers import blk, body_text_nodes, do_render, docx_all_text, docx_of, docx_xml, snap

pytestmark = pytest.mark.db
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
STYLE = {"font": "liberation-serif", "font_size_pt": 11, "primary_color": "#123456", "accent_color": "#abcdef",
         "header_text": "Example Assistant assessment", "footer_text": "For internal review", "marking": "internal"}
MARKDOWN = ("Intro with **bold words** and *slanted words* and `code_bit`.\n\n"
            "- first bullet\n- second bullet\n\n1. first step\n2. second step\n\n"
            "| Name | Value |\n|---|---|\n| alpha | 1 |\n| beta | 2 |\n\n"
            "See [the example site](https://example.org).")


def layout():
    return [
        blk("cover", report_title="Example Assistant report"),
        blk("chapter", title="Evidence", intro="Why the evidence matters."),
        blk("free_text", title="Method", format="markdown", text=MARKDOWN),
        blk("ai_card", commentary="Checked by the DPO."),
        blk("test_results", page_break_before=True),
        blk("chart", dataset="checklist_scores"),
        blk("appendix"),
        blk("free_text", title="Glossary", text="Plain glossary text."),
    ]


@pytest.fixture(scope="module")
def rendered(bed, sources):
    """One docx and one preview of the same layout (module scope: rendering is the slow part)."""
    from conftest import FIXED_NOW, RecordingImages

    try:
        build_registry = need("report_renderer.registry", "build_registry")
        Deps = need("report_renderer.context", "Deps")
    except BaseException as exc:  # noqa: BLE001
        return {"error": str(exc)}
    blocks = layout()
    try:
        deps = Deps(registry=build_registry(), sources=sources, images=RecordingImages(), clock=lambda: FIXED_NOW,
                    chart_data=None)
    except BaseException as exc:  # noqa: BLE001 - reported in each test body
        return {"error": f"missing feature: renderer deps ({type(exc).__name__}: {exc})"}
    common = dict(project="M", system="M_V2", style=STYLE, language="en",
                  document={"id": "3c9d1f2e-7a6b-4c5d-8e9f-0a1b2c3d4e5f", "toc": "on", "numbering": False})
    s = snap(blocks, mode="docx", **common)
    s["layout"]["name"] = "Example Assistant layout"
    try:
        result = do_render(deps, s)
        preview = do_render(deps, dict(s, mode="preview"))
    except BaseException as exc:  # noqa: BLE001 - a failed fixture would be an ERROR, not a FAILED test
        return {"error": f"missing feature: mode=docx rendering ({type(exc).__name__}: {str(exc)[:300]})"}
    return {"result": result, "preview": preview, "blocks": blocks}


def got(rendered, key="result"):
    if "error" in rendered:
        pytest.fail(rendered["error"], pytrace=False)
    return rendered[key]


def paragraphs(document):
    return [(p.style.name, " ".join(p.text.split())) for p in document.paragraphs]


def test_docx_mode_returns_the_document_and_its_hashes(rendered):
    r = got(rendered)
    assert {"docx_base64", "sha256", "fingerprint", "block_statuses"} <= set(r), sorted(r)
    import base64
    import hashlib

    raw = base64.b64decode(r["docx_base64"])
    assert raw[:2] == b"PK" and hashlib.sha256(raw).hexdigest() == r["sha256"]


def test_headings_map_to_word_styles(rendered):
    paras = paragraphs(docx_of(got(rendered)))
    assert ("Title", "Example Assistant report") in paras
    assert ("Heading 1", "Evidence") in paras
    assert ("Heading 2", "Method") in paras
    assert any(style == "Normal" and "Why the evidence matters." in t for style, t in paras)


def test_inline_formatting_becomes_runs(rendered):
    d = docx_of(got(rendered))
    runs = [r for p in d.paragraphs for r in p.runs]
    assert any(r.bold and "bold words" in r.text for r in runs)
    assert any(r.italic and "slanted words" in r.text for r in runs)
    mono = [r for r in runs if "code_bit" in r.text]
    assert mono and (mono[0].font.name or "").lower() in ("courier new", "consolas", "dejavu sans mono",
                                                          "liberation mono", "courier")


def test_links_become_hyperlinks(rendered):
    d = docx_of(got(rendered))
    rels = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "https://example.org" in rels
    assert "<w:hyperlink" in docx_xml(got(rendered), "word/document.xml")


def test_lists_become_bullets_and_numbers(rendered):
    d = docx_of(got(rendered))
    listed = {" ".join(p.text.split()): p for p in d.paragraphs
              if p._p.pPr is not None and p._p.pPr.numPr is not None or "List" in p.style.name}
    for item in ("first bullet", "second bullet", "first step", "second step"):
        assert item in listed, f"{item} is not a list paragraph"
    assert listed["first bullet"].style.name != listed["first step"].style.name or \
        listed["first bullet"]._p.pPr.numPr.numId.val != listed["first step"]._p.pPr.numPr.numId.val


def test_tables_have_a_bold_repeated_header(rendered):
    d = docx_of(got(rendered))
    tables = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:2] == ["Name", "Value"]]
    assert tables, "the markdown table is a Word table"
    header = tables[0].rows[0]
    assert header._tr.trPr is not None and header._tr.trPr.find(f"{W}tblHeader") is not None
    assert all(r.bold for c in header.cells for p in c.paragraphs for r in p.runs if r.text.strip())
    assert [c.text.strip() for c in tables[0].rows[2].cells] == ["beta", "2"]


def test_charts_become_pictures_or_their_table(rendered):
    d = docx_of(got(rendered))
    all_text = docx_all_text(d)
    pictures = len(d.inline_shapes)
    assert pictures >= 1 or "Chart shown as a table." in all_text
    assert "Example Assistant transparency" in all_text          # the chart's data table (show_table default true)


def test_a_page_break_before_is_a_page_break(rendered):
    xml = docx_xml(got(rendered), "word/document.xml")
    assert re.search(r'<w:br [^>]*w:type="page"', xml) or "<w:pageBreakBefore" in xml


def test_commentary_is_a_shaded_or_bordered_paragraph(rendered):
    d = docx_of(got(rendered))
    p = next(p for p in d.paragraphs if "Checked by the DPO." in p.text)
    ppr = p._p.pPr
    assert ppr is not None and (ppr.find(f"{W}shd") is not None or ppr.find(f"{W}pBdr") is not None)


def test_no_text_of_the_html_body_is_lost(rendered):
    all_text = docx_all_text(docx_of(got(rendered)))
    # The cover says "Generated by: preview" in a preview and "Generated by {requester}" in a document
    # (the DOCX is built from the PDF's HTML): that one line differs by design.
    nodes = [t for t in body_text_nodes(got(rendered, "preview")["html"]) if t != "Generated by: preview"]
    missing = [t for t in nodes if t not in all_text]
    assert not missing, f"texts missing in the DOCX: {missing[:10]}"


def test_header_and_footer_carry_the_pdf_texts_and_page_fields(rendered):
    d = docx_of(got(rendered))
    sec = d.sections[0]
    header = " ".join(p.text for p in sec.header.paragraphs)
    assert "Example Assistant assessment" in header and "INTERNAL" in header
    assert "For internal review" in " ".join(p.text for p in sec.footer.paragraphs)
    footer_xml = docx_xml(got(rendered), r"word/footer\d*\.xml")
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml


def test_the_toc_is_a_word_field_updated_on_open(rendered):
    xml = docx_xml(got(rendered), "word/document.xml")
    assert re.search(r'TOC \\o "1-2"', xml), "a TOC field with levels 1 and 2"
    assert "Method" in xml
    settings = docx_xml(got(rendered), "word/settings.xml")
    assert re.search(r'<w:updateFields w:val="(true|1)"', settings)


def test_fonts_follow_the_template(rendered):
    d = docx_of(got(rendered))
    normal = d.styles["Normal"].font
    assert normal.name == "Liberation Serif"
    assert normal.size is not None and normal.size.pt == 11
    colour = d.styles["Heading 2"].font.color.rgb
    assert colour is not None and str(colour).upper() == "123456"


def test_document_properties(rendered):
    props = docx_of(got(rendered)).core_properties
    assert props.title == "Example Assistant layout"
    assert props.subject == "Mike project"
    assert props.language == "en"
    assert "alice" not in (props.author or "") and "alice" not in (props.last_modified_by or "")


# Characters XML 1.0 forbids never break the Word file.

#: one sample of every class of character XML 1.0 forbids (NUL, the other C0 controls except tab, LF and CR
#: (including VT and FF), a lone high and a lone low surrogate, U+FFFE and U+FFFF)
FORBIDDEN = {"nul": "\x00", "soh": "\x01", "bs": "\x08", "vt": "\x0b", "ff": "\x0c", "so": "\x0e", "esc": "\x1b",
             "us": "\x1f", "high_surrogate": "\ud800", "low_surrogate": "\udfff", "fffe": "￾", "ffff": "￿"}


def _xml_forbidden(text: str) -> list[str]:
    return [c for c in text if (ord(c) < 0x20 and c not in "\t\n\r") or 0xD800 <= ord(c) <= 0xDFFF
            or c in "￾￿"]


def _chart_svg(label: str) -> str:
    from aisc_report_plugin_interface import charts

    return str(charts.svg({"id": "c", "title": "Scores " + label, "kind": "bar", "value_format": "number",
                           "axis_min": 0, "axis_max": 5, "series": [{"label": "L" + label + "x", "value": 3}]},
                          {"primary": "#000fdf", "text": "#111111", "grid": "#dddddd"}))


@pytest.mark.parametrize("name", sorted(FORBIDDEN))
def test_html_to_docx_makes_every_text_xml_safe(name):
    """Every place the DOCX writer puts text: title, headings, paragraphs, commentary, list items, table
    cells, definition lists, links, the TOC, a picture's alt text, chart labels and the document properties."""
    from html import escape as esc

    html_to_docx = need("report_renderer.docx", "html_to_docx")
    c = FORBIDDEN[name]
    e = esc(c)
    html = (f'<html><body><nav id="toc"><h2>Con{e}tents</h2><ol><li><a href="#x">Ent{e}ry</a></li></ol></nav>'
            f'<h1>Ti{e}tle</h1><h1 class="chapter-title">Chap{e}ter</h1><h2>Sec{e}tion</h2>'
            f'<p>Free{e} text <strong>bo{e}ld</strong> <a href="https://e.org/">li{e}nk</a></p>'
            f'<div class="commentary"><p>Com{e}mentary</p></div>'
            f'<ul><li>It{e}em</li></ul>'
            f'<table><thead><tr><th>Hea{e}d</th></tr></thead><tbody><tr><td>Ce{e}ll</td></tr></tbody></table>'
            f'<dl><dt>Te{e}rm</dt><dd>Defi{e}nition</dd></dl>'
            f'<img alt="Al{e}t" src="missing.png">'
            f'<figure>{_chart_svg(c)}<figcaption>Cap{e}tion</figcaption></figure>'
            f'</body></html>')
    raw = html_to_docx(html, title="Doc" + c, subject="Sub" + c, language="en", look=None, t=lambda s, **k: s,
                       margins={"top_left": "M" + c})
    import io
    import zipfile

    import docx

    document = docx.Document(io.BytesIO(raw))
    body = docx_all_text(document)
    for word in ("Con", "tents", "Ent", "ry", "Ti", "tle", "Chap", "ter", "Sec", "tion", "Free", "bo", "ld", "li",
                 "nk", "Com", "mentary", "It", "em", "Hea", "Ce", "ll", "Te", "rm", "Defi", "nition", "Al", "Cap"):
        assert word in body, f"{word!r} lost from the Word file"
    assert not _xml_forbidden(body)
    assert len(document.inline_shapes) == 1, "the chart with the character in its labels is still a picture"
    assert document.core_properties.title.startswith("Doc") and document.core_properties.subject.startswith("Sub")
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        for n in z.namelist():
            if n.endswith(".xml"):
                assert not _xml_forbidden(z.read(n).decode("utf-8")), n


@pytest.mark.parametrize("name", sorted(set(FORBIDDEN) - {"high_surrogate", "low_surrogate"}))
def test_a_docx_render_with_forbidden_characters_in_the_texts_succeeds(deps, name):
    """The verifier's case through the whole renderer: free text, markdown table cells, commentary, cover
    title and subtitle, chapter title and intro, block titles. (Lone surrogates are refused earlier, when the
    snapshot is checked.)"""
    c = FORBIDDEN[name]
    blocks = [blk("cover", report_title="Rep" + c + "ort", subtitle="Sub" + c + "title"),
              blk("chapter", title="Chap" + c + "ter", intro="In" + c + "tro"),
              blk("free_text", title="Fr" + c + "ee", format="markdown",
                  text="a" + c + "b\n\n| Na" + c + "me | V |\n|---|---|\n| ce" + c + "ll | 1 |\n"),
              blk("free_text", title="Pl" + c + "ain", text="pl" + c + "ain text"),
              blk("ai_card", commentary="c" + c + "omm")]
    d = docx_of(do_render(deps, snap(blocks, mode="docx")))
    body = docx_all_text(d)
    for word in ("Rep", "ort", "Sub", "title", "Chap", "ter", "In", "tro", "Fr", "ee", "Na", "me", "ce", "ll", "Pl",
                 "ain", "pl", "omm"):
        assert word in body, f"{word!r} lost from the Word file"
    assert not _xml_forbidden(body)


@pytest.mark.parametrize("ref", ["&#1;", "&#x1b;", "&#xFFFE;", "&#xffff;", "&#x0c;"])
def test_character_references_to_forbidden_characters_are_safe_too(ref):
    """Markdown passes `&#1;` through as HTML, and lxml turns it into the character itself; that text also
    goes through the nested-block paths (a list item or a table cell holding a paragraph)."""
    import io

    import docx

    html_to_docx = need("report_renderer.docx", "html_to_docx")
    html = (f'<html><body><p>Pa{ref}ra</p><ul><li><p>Ne{ref}sted</p> tail</li></ul>'
            f'<table><tr><td><p>Blo{ref}ck</p></td></tr></table><h2>He{ref}ad</h2>'
            f'<p><a href="https://e.org/">Li{ref}nk</a></p><nav id="toc"><ol><li><a href="#a">To{ref}c</a></li></ol></nav>'
            f'</body></html>')
    body = docx_all_text(docx.Document(io.BytesIO(
        html_to_docx(html, title="t", subject="s", language="en", look=None, t=lambda s, **k: s))))
    for word in ("Pa", "ra", "Ne", "sted", "Blo", "ck", "He", "ad", "Li", "nk", "To"):
        assert word in body
    assert not _xml_forbidden(body)


def test_a_character_reference_in_markdown_free_text_renders_as_docx(deps):
    text = "a&#1;b\n\n- it&#x1b;em\n\n| h&#1; | v |\n|---|---|\n| c&#xFFFE; | 1 |\n"
    d = docx_of(do_render(deps, snap([blk("free_text", format="markdown", text=text)], mode="docx")))
    body = docx_all_text(d)
    assert "it" in body and "em" in body and not _xml_forbidden(body)


def test_a_chart_inside_a_list_item_does_not_shift_the_later_charts():
    """SVGs are taken in document order; one inside a list item (or a table cell with blocks) was
    flattened to text without taking its SVG, so every later chart showed the one before it."""
    import io
    import zipfile

    import cairosvg

    from report_renderer.docx import html_to_docx

    first = '<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><rect width="10" height="10" fill="#ff0000"/></svg>'
    second = '<svg xmlns="http://www.w3.org/2000/svg" width="12" height="12"><rect width="12" height="12" fill="#0000ff"/></svg>'
    html = (f"<html><body><ul><li><div>in a list</div>{first}</li></ul>"
            f"<figure>{second}</figure></body></html>")
    data = html_to_docx(html, title="t", subject="s", language="en", look=None, t=lambda s, **k: s)
    media = [zipfile.ZipFile(io.BytesIO(data)).read(n) for n in zipfile.ZipFile(io.BytesIO(data)).namelist()
             if n.startswith("word/media/")]
    assert cairosvg.svg2png(bytestring=second.encode(), scale=2) in media
    assert cairosvg.svg2png(bytestring=first.encode(), scale=2) not in media


def test_an_svg_picture_is_rasterised_at_a_bounded_size(monkeypatch):
    """The Word export drew an SVG <img> at scale=2 of its declared
    size, so a 16000 px logo asked cairo for a 32000 x 32000 surface (3.8 GiB). The picture (and a chart)
    is now drawn with its longer side at most SVG_RASTER_MAX_SIDE pixels."""
    import base64

    import cairosvg

    from report_renderer import style
    from report_renderer.docx import html_to_docx

    asked = []
    real = cairosvg.svg2png

    def recording(*, bytestring, **kw):
        size = (16000, 16000) if b'width="16000"' in bytestring else (10, 9000)
        if "output_width" in kw or "output_height" in kw:
            asked.append((kw.get("output_width"), kw.get("output_height")))
        else:
            asked.append((size[0] * kw.get("scale", 1), size[1] * kw.get("scale", 1)))
        return real(bytestring=b'<svg xmlns="http://www.w3.org/2000/svg" width="1" height="1"/>')

    monkeypatch.setattr(cairosvg, "svg2png", recording)
    big = '<svg xmlns="http://www.w3.org/2000/svg" width="16000" height="16000"><rect width="10" height="10"/></svg>'
    tall = '<svg xmlns="http://www.w3.org/2000/svg" width="10" height="9000"><rect width="10" height="10"/></svg>'
    logo = base64.b64encode(big.encode()).decode()
    html = (f'<html><body><section class="cover"><img src="data:image/svg+xml;base64,{logo}" alt="logo"></section>'
            f"<figure>{tall}</figure></body></html>")
    html_to_docx(html, title="t", subject="s", language="en", look=None, t=lambda s, **k: s)
    assert len(asked) == 2
    for w, h in asked:
        assert w * h <= style.LOGO_MAX_PIXELS
        assert max(w, h) <= style.SVG_RASTER_MAX_SIDE


def test_a_normal_svg_picture_keeps_twice_its_size():
    import base64
    import io
    import zipfile

    import cairosvg

    from report_renderer.docx import html_to_docx

    svg = '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="80"><rect width="200" height="80"/></svg>'
    html = f'<html><body><img src="data:image/svg+xml;base64,{base64.b64encode(svg.encode()).decode()}"></body></html>'
    data = html_to_docx(html, title="t", subject="s", language="en", look=None, t=lambda s, **k: s)
    z = zipfile.ZipFile(io.BytesIO(data))
    media = [z.read(n) for n in z.namelist() if n.startswith("word/media/")]
    assert media == [cairosvg.svg2png(bytestring=svg.encode(), scale=2)]
